"""CA-R-01 a CA-R-08 dos rankings técnico, comportamental e final."""

from __future__ import annotations

import json
from io import BytesIO

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine

from avaliacao_ia.exportar import COLUNAS as COLUNAS_RESULTADO
from avaliacao_ia.exportar import juntar_perguntas, tabela_exportacao
from avaliacao_ia.ranking import atribuir_posicoes, media_meio_para_cima, notas_para_gravar
from avaliacao_ia.validacao import ErroValidacaoJson, validar_saida_modelo
from triagem.exportar import COLUNAS as COLUNAS_TRIAGEM

PERGUNTA = "Como você atende o cliente?"
RESPOSTA = "eu explico de outro jeito"
JD = "A vaga pede atendimento consultivo e clareza na escrita."
CARGO = "Analista de dados"


@pytest.fixture(autouse=True)
def sem_chamada_externa(monkeypatch):
    class ClienteHttpBloqueado(httpx.Client):
        def post(self, *args, **kwargs):
            raise AssertionError("chamada externa")

    monkeypatch.setattr("avaliacao_ia.cliente.httpx.Client", ClienteHttpBloqueado)

    def completar_proibido(self, mensagens):
        raise AssertionError("chamada externa")

    monkeypatch.setattr("avaliacao_ia.cliente.ClienteOpenAICompat.completar", completar_proibido)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("SCREENER_DB_PATH", str(tmp_path / "screener.sqlite"))
    import persistencia.db as banco

    banco._engine = None
    banco.init_db()
    from main import app

    with TestClient(app) as c:
        yield c


def _corpo(
    tecnica: int,
    comportamental: int,
    resultado: str = "atende",
    extra: dict | None = None,
    pergunta: str = PERGUNTA,
    resposta: str = RESPOSTA,
) -> str:
    corpo = {
        "perfil": "A resposta foi lida contra a pergunta e a job description.",
        "por_pergunta": [
            {
                "perguntaId": "p1",
                "pergunta": pergunta,
                "resultado": resultado,
                "justificativa": "A resposta foi confrontada com a pergunta.",
            }
        ],
        "trechos": [
            {
                "perguntaId": "p1",
                "trecho": "" if resultado == "sem_informacao" else resposta,
            }
        ],
        "nota_tecnica": tecnica,
        "nota_comportamental": comportamental,
        "justificativa_nota_tecnica": "A resposta cobre o pedido da vaga. Isso cabe na faixa da rubrica técnica.",
        "justificativa_nota_comportamental": "O texto é claro e organizado. Isso cabe na faixa da escrita.",
    }
    if extra:
        corpo.update(extra)
    return json.dumps(corpo)


def _csv(pessoas: list[tuple[str, str]]) -> str:
    linhas = [f"Nome,Cargo,Idade,Escolaridade,Telefone,Email,Busca LKD,{PERGUNTA}"]
    for indice, (nome, cargo) in enumerate(pessoas, start=1):
        email = f"{nome.split()[0].lower()}@exemplo.com"
        linhas.append(
            f"{nome},{cargo},30,Bacharelado,1199999000{indice},{email},busca-{indice},{RESPOSTA}"
        )
    return "\n".join(linhas) + "\n"


def _preparar(client: TestClient, pessoas: list[tuple[str, str]] | None = None) -> tuple[str, dict[str, dict]]:
    if pessoas is None:
        pessoas = [
            ("Fulana Costa", CARGO),
            ("Beltrana Costa", CARGO),
            ("Ciclano Souza", "Operador de caixa"),
        ]
    criado = client.post("/api/projetos", json={"nome": "Vaga"})
    assert criado.status_code == 200, criado.text
    pid = criado.json()["projetoId"]
    criterios = client.put(
        f"/api/projetos/{pid}/criterios-deterministicos",
        json={
            "cargos": {"ativo": True, "aceitos": [CARGO]},
            "idade": {"ativo": False, "minima": None, "maxima": None},
            "escolaridade": {"ativo": False, "configuracao": {"modelo": "lista_valores_aceitos", "aceitos": []}},
        },
    )
    assert criterios.status_code == 200, criterios.text
    importado = client.post(
        f"/api/projetos/{pid}/importacao/confirmar",
        files={"arquivo": ("lote.csv", _csv(pessoas).encode("utf-8"), "text/csv")},
        data={
            "encoding": "utf-8",
            "mapeamento_json": json.dumps(
                {
                    "ultimoCargo": "Cargo",
                    "idade": "Idade",
                    "escolaridade": "Escolaridade",
                    "nomeExibicao": "Nome",
                    "respostasChatbot": [PERGUNTA],
                    "cidade": None,
                    "statusLegado": None,
                }
            ),
        },
    )
    assert importado.status_code == 200, importado.text
    assert client.post(f"/api/projetos/{pid}/triagem/executar").status_code == 200
    assert client.put(f"/api/projetos/{pid}/job-description", json={"texto": JD}).status_code == 200
    prompt = client.get(f"/api/projetos/{pid}/prompt").json()
    salvo = client.put(f"/api/projetos/{pid}/prompt", json={"template": prompt["templateInicial"]})
    assert salvo.status_code == 200, salvo.text
    linhas = client.get(f"/api/projetos/{pid}/triagem").json()["linhas"]
    return pid, {linha["nomeExibicao"]: linha for linha in linhas}


class _Fila:
    def __init__(self, corpos: list[str]) -> None:
        self.corpos = corpos
        self.chamadas = 0

    def completar(self, mensagens):
        corpo = self.corpos[self.chamadas]
        self.chamadas += 1
        return {
            "conteudo": corpo,
            "uso": {"entrada": 4, "saida": 3, "total": 7, "nota": None},
            "tentativasRede": 1,
        }


def _ligar(monkeypatch, corpos: list[str]) -> _Fila:
    fila = _Fila(corpos)
    monkeypatch.setattr("avaliacao_ia.servico.config.llm_configurada", lambda: True)
    monkeypatch.setattr("avaliacao_ia.servico.cliente_padrao", lambda: fila)
    return fila


def _avaliar(client: TestClient, pid: str, ids: list[str]):
    disparo = client.post(
        f"/api/projetos/{pid}/avaliacao-ia/executar",
        json={"candidatoIds": ids, "confirmar": True},
    )
    assert disparo.status_code == 200, disparo.text
    return disparo.json()


def _ler(client: TestClient, pid: str) -> list[dict]:
    corpo = client.get(f"/api/projetos/{pid}/resultado")
    assert corpo.status_code == 200, corpo.text
    return corpo.json()["linhas"]


def _por_nome(linhas: list[dict]) -> dict[str, dict]:
    return {linha["nomeExibicao"]: linha for linha in linhas}


def test_ca_r_01_media_meio_para_cima():
    assert media_meio_para_cima(78, 70) == 74
    assert media_meio_para_cima(78, 71) == 75
    assert media_meio_para_cima(0, 1) == 1
    assert media_meio_para_cima(0, 0) == 0
    assert round((78 + 71) / 2) == 74


def test_ca_r_01_media_gravada_na_leitura(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    beltrana = pessoas["Beltrana Costa"]
    _ligar(
        monkeypatch,
        [_corpo(78, 70), _corpo(78, 71)],
    )
    _avaliar(client, pid, [fulana["candidatoId"], beltrana["candidatoId"]])
    linhas = _por_nome(_ler(client, pid))
    assert linhas["Fulana Costa"]["avaliacao"]["notaTecnica"] == 78
    assert linhas["Fulana Costa"]["avaliacao"]["notaComportamental"] == 70
    assert linhas["Fulana Costa"]["ranking"]["notaFinal"] == 74
    assert linhas["Beltrana Costa"]["ranking"]["notaFinal"] == 75
    assert "notaFinal" not in linhas["Fulana Costa"]["avaliacao"]


def test_ca_r_02_chave_ranking_nao_gera_posicao(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    fila = _ligar(monkeypatch, [_corpo(80, 70, extra={"ranking": 1}), _corpo(80, 70, extra={"ranking": 1})])
    corpo = _avaliar(client, pid, [fulana["candidatoId"]])
    assert fila.chamadas == 2
    assert corpo["resultados"][0]["estado"] == "falha_avaliacao"
    linha = _por_nome(_ler(client, pid))["Fulana Costa"]
    assert linha["resultadoExibicao"] == "falha_avaliacao"
    assert linha["avaliacao"]["notaTecnica"] is None
    assert linha["avaliacao"]["notaComportamental"] is None
    assert linha["ranking"]["notaFinal"] is None
    assert linha["ranking"]["posicaoFinal"] is None
    assert linha["ranking"]["posicaoTecnica"] is None
    assert linha["ranking"]["posicaoComportamental"] is None
    assert linha["ranking"]["totalComNota"] == 0


def test_chave_nota_final_continua_sendo_falha_de_schema():
    respostas = [{"colunaMapeada": PERGUNTA, "texto": RESPOSTA}]
    with pytest.raises(ErroValidacaoJson):
        validar_saida_modelo(_corpo(78, 70, extra={"nota_final": 74}), respostas)
    with pytest.raises(ErroValidacaoJson):
        validar_saida_modelo(_corpo(78, 70, extra={"ranking": 1}), respostas)


def test_ca_r_03_ordem_tecnica_com_desempate(client: TestClient, monkeypatch):
    pessoas = [("Cida", CARGO)]
    pessoas.append(("Ana", CARGO))
    pessoas.extend((f"Fora {indice}", "Operador de caixa") for indice in range(3, 9))
    pessoas.append(("Bia", CARGO))
    pid, por_nome = _preparar(client, pessoas)
    assert por_nome["Cida"]["linhaOrigem"] == 1
    assert por_nome["Ana"]["linhaOrigem"] == 2
    assert por_nome["Bia"]["linhaOrigem"] == 9
    ids = [por_nome["Ana"]["candidatoId"], por_nome["Bia"]["candidatoId"], por_nome["Cida"]["candidatoId"]]
    _ligar(monkeypatch, [_corpo(80, 40), _corpo(80, 90), _corpo(70, 99)])
    _avaliar(client, pid, ids)
    linhas = _por_nome(_ler(client, pid))
    ordem_tecnica = sorted(
        ("Ana", "Bia", "Cida"),
        key=lambda nome: linhas[nome]["ranking"]["posicaoTecnica"],
    )
    assert ordem_tecnica == ["Bia", "Ana", "Cida"]
    assert [linhas[nome]["ranking"]["posicaoTecnica"] for nome in ordem_tecnica] == [1, 2, 3]
    assert linhas["Fora 3"]["ranking"]["posicaoTecnica"] is None
    assert linhas["Bia"]["ranking"]["totalComNota"] == 3
    assert linhas["Fora 3"]["ranking"]["totalComNota"] == 3


def test_desempate_por_linha_quando_as_notas_empatam():
    entradas = [
        {"candidatoId": "tarde", "linhaOrigem": 9, "notaTecnica": 50, "notaComportamental": 50, "notaFinal": 50},
        {"candidatoId": "cedo", "linhaOrigem": 1, "notaTecnica": 50, "notaComportamental": 50, "notaFinal": 50},
    ]
    posicoes = atribuir_posicoes(entradas)
    assert posicoes["cedo"]["posicaoTecnica"] == 1
    assert posicoes["tarde"]["posicaoTecnica"] == 2
    assert posicoes["cedo"]["posicaoComportamental"] == 1
    assert posicoes["cedo"]["posicaoFinal"] == 1


def test_ca_r_04_inelegivel_fora_do_denominador(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    beltrana = pessoas["Beltrana Costa"]
    _ligar(monkeypatch, [_corpo(80, 60), _corpo(40, 40)])
    _avaliar(client, pid, [fulana["candidatoId"], beltrana["candidatoId"]])
    linhas = _por_nome(_ler(client, pid))
    ciclano = linhas["Ciclano Souza"]
    assert ciclano["resultadoExibicao"] == "inelegivel"
    assert ciclano["ranking"]["posicaoFinal"] is None
    assert ciclano["ranking"]["posicaoTecnica"] is None
    assert ciclano["ranking"]["posicaoComportamental"] is None
    assert ciclano["ranking"]["notaFinal"] is None
    assert ciclano["ranking"]["totalComNota"] == 2
    assert linhas["Fulana Costa"]["ranking"]["totalComNota"] == 2
    posicoes = {
        linhas["Fulana Costa"]["ranking"]["posicaoFinal"],
        linhas["Beltrana Costa"]["ranking"]["posicaoFinal"],
    }
    assert posicoes == {1, 2}


def test_ca_r_05_zeramento_tecnico_quando_tudo_e_sem_informacao(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    _ligar(monkeypatch, [_corpo(90, 40, resultado="sem_informacao")])
    _avaliar(client, pid, [fulana["candidatoId"]])
    linha = _por_nome(_ler(client, pid))["Fulana Costa"]
    assert linha["avaliacao"]["notaTecnica"] == 0
    assert linha["avaliacao"]["notaComportamental"] == 40
    assert linha["ranking"]["notaFinal"] == 20
    assert linha["ranking"]["notaTecnicaZeradaPorFaltaDeInformacao"] is True
    assert linha["ranking"]["posicaoTecnica"] == 1
    assert linha["ranking"]["totalComNota"] == 1


def test_zeramento_so_quando_todas_as_perguntas_sao_sem_informacao():
    parcial = notas_para_gravar(
        90,
        40,
        [{"resultado": "sem_informacao"}, {"resultado": "atende"}],
    )
    assert parcial["notaTecnica"] == 90
    assert parcial["notaComportamental"] == 40
    assert parcial["notaTecnicaZeradaPorFaltaDeInformacao"] is False
    assert parcial["notaFinal"] == media_meio_para_cima(90, 40)

    total = notas_para_gravar(90, 40, [{"resultado": "sem_informacao"}])
    assert total["notaTecnica"] == 0
    assert total["notaFinal"] == 20
    assert total["notaTecnicaZeradaPorFaltaDeInformacao"] is True


def test_ca_r_06_eixos_nao_se_copiam(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    _ligar(monkeypatch, [_corpo(12, 91)])
    _avaliar(client, pid, [fulana["candidatoId"]])
    linha = _por_nome(_ler(client, pid))["Fulana Costa"]
    assert linha["avaliacao"]["notaTecnica"] == 12
    assert linha["avaliacao"]["notaComportamental"] == 91
    assert linha["ranking"]["notaFinal"] == media_meio_para_cima(12, 91)
    assert linha["ranking"]["notaTecnicaZeradaPorFaltaDeInformacao"] is False
    assert linha["avaliacao"]["notaTecnica"] != linha["avaliacao"]["notaComportamental"]


def test_ca_r_07_reexecucao_recalcula_o_lote(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    beltrana = pessoas["Beltrana Costa"]
    fila = _ligar(monkeypatch, [_corpo(80, 80), _corpo(40, 40), _corpo(96, 96)])
    _avaliar(client, pid, [fulana["candidatoId"], beltrana["candidatoId"]])
    antes = _por_nome(_ler(client, pid))
    assert antes["Fulana Costa"]["ranking"]["notaFinal"] == 80
    assert antes["Beltrana Costa"]["ranking"]["notaFinal"] == 40
    assert antes["Fulana Costa"]["ranking"]["posicaoFinal"] == 1
    assert antes["Beltrana Costa"]["ranking"]["posicaoFinal"] == 2
    _avaliar(client, pid, [beltrana["candidatoId"]])
    assert fila.chamadas == 3
    depois = _por_nome(_ler(client, pid))
    assert depois["Beltrana Costa"]["ranking"]["notaFinal"] == 96
    assert depois["Beltrana Costa"]["ranking"]["posicaoFinal"] == 1
    assert depois["Fulana Costa"]["ranking"]["posicaoFinal"] == 2
    assert depois["Fulana Costa"]["ranking"]["notaFinal"] == 80
    assert depois["Ciclano Souza"]["ranking"]["totalComNota"] == 2


def test_ca_r_08_avaliacao_antiga_nao_ganha_nota(client: TestClient):
    from persistencia.db import SessionLocal
    from persistencia.models import AvaliacaoIa

    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    db = SessionLocal()
    try:
        db.add(
            AvaliacaoIa(
                candidato_id=fulana["candidatoId"],
                projeto_id=pid,
                prompt_version_id="prm_antigo",
                criterios_qualitativos_versao_id="cq_antigo",
                perfil="Perfil do schema antigo.",
                compatibilidade="parcial",
                justificativas_por_criterio=[
                    {
                        "criterioId": "cq_01",
                        "nome": "Escuta",
                        "resultado": "atende_parcialmente",
                        "justificativa": "Cobre parte do critério.",
                        "trecho": RESPOSTA,
                    }
                ],
                por_pergunta=None,
                trechos=None,
                nota_tecnica=None,
                nota_comportamental=None,
                nota_final=None,
                uso_token={"entrada": 1, "saida": 1, "total": 2},
                estado="sucesso",
                erro=None,
            )
        )
        db.commit()
    finally:
        db.close()
    linha = _por_nome(_ler(client, pid))["Fulana Costa"]
    assert linha["resultadoExibicao"] == "avaliado"
    assert linha["avaliacao"]["compatibilidade"] == "parcial"
    assert linha["avaliacao"]["notaTecnica"] is None
    assert linha["avaliacao"]["notaComportamental"] is None
    assert linha["avaliacao"]["porPergunta"] is None
    assert linha["ranking"]["notaFinal"] is None
    assert linha["ranking"]["posicaoFinal"] is None
    assert linha["ranking"]["posicaoTecnica"] is None
    assert linha["ranking"]["posicaoComportamental"] is None
    assert linha["ranking"]["totalComNota"] == 0
    assert linha["ranking"]["notaTecnicaZeradaPorFaltaDeInformacao"] is False


def test_rotulos_das_perguntas_na_exportacao():
    texto = juntar_perguntas(
        [
            {"pergunta": "P1", "resultado": "atende"},
            {"pergunta": "P2", "resultado": "atende_parcialmente"},
            {"pergunta": "P3", "resultado": "nao_atende"},
            {"pergunta": "P4", "resultado": "sem_informacao"},
        ]
    )
    assert texto == "P1: atende | P2: atende parcialmente | P3: não atende | P4: sem informação"


def test_justificativas_em_uma_coluna_por_pergunta():
    por_pergunta = [
        {"pergunta": "P1", "resultado": "atende", "justificativa": "J1"},
        {"pergunta": "P2", "resultado": "atende", "justificativa": "J2"},
        {"pergunta": "P3", "resultado": "nao_atende", "justificativa": "J3"},
    ]
    linhas = [
        {
            "candidatoId": "c1",
            "linhaOrigem": 1,
            "ranking": {
                "posicaoFinal": 1,
                "posicaoTecnica": 1,
                "posicaoComportamental": 1,
                "notaFinal": 80,
            },
            "avaliacao": {
                "notaTecnica": 80,
                "notaComportamental": 80,
                "perfil": "Resumo.",
                "porPergunta": por_pergunta,
            },
        },
        {
            "candidatoId": "c2",
            "linhaOrigem": 2,
            "ranking": {"posicaoFinal": None},
            "avaliacao": None,
        },
    ]
    cabecalho, tabela = tabela_exportacao(linhas, {"c1": {"nome": "Ana"}})
    assert cabecalho[-3:] == [
        "Justificativa pergunta 1",
        "Justificativa pergunta 2",
        "Justificativa pergunta 3",
    ]
    assert tabela[0][-3:] == ["J1", "J2", "J3"]
    assert tabela[1][-3:] == [None, None, None]


def test_excel_do_resultado_nao_substitui_o_da_triagem(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    _ligar(monkeypatch, [_corpo(78, 70)])
    _avaliar(client, pid, [fulana["candidatoId"]])

    resultado = client.get(f"/api/projetos/{pid}/resultado/exportar")
    assert resultado.status_code == 200, resultado.text
    assert "spreadsheetml" in resultado.headers["content-type"]
    livro = load_workbook(BytesIO(resultado.content))
    aba = livro.active
    assert aba.title == "Resultado"
    assert [celula.value for celula in aba[1]] == [
        *COLUNAS_RESULTADO,
        "Justificativa pergunta 1",
    ]
    por_linha = {linha[0]: linha for linha in aba.iter_rows(min_row=2, values_only=True)}
    avaliada = por_linha[1]
    assert avaliada[1] == "Fulana Costa"
    assert avaliada[2] == "fulana@exemplo.com"
    assert avaliada[3] == "11999990001"
    assert avaliada[4] == 78
    assert avaliada[5] == "A resposta cobre o pedido da vaga. Isso cabe na faixa da rubrica técnica."
    assert avaliada[6] == 1
    assert avaliada[7] == 70
    assert avaliada[8] == "O texto é claro e organizado. Isso cabe na faixa da escrita."
    assert avaliada[9] == 1
    assert avaliada[10] == 74
    assert avaliada[11] == 1
    assert avaliada[12] == f"{PERGUNTA}: atende"
    assert avaliada[13] == "A resposta foi lida contra a pergunta e a job description."
    assert avaliada[14] == "A resposta foi confrontada com a pergunta."
    sem_nota = por_linha[3]
    assert sem_nota[1] == "Ciclano Souza"
    assert sem_nota[4:12] == (None, None, None, None, None, None, None, None)
    assert sem_nota[12] in (None, "")
    assert sem_nota[13] in (None, "")
    assert sem_nota[14] in (None, "")

    so_nota = client.get(f"/api/projetos/{pid}/resultado/exportar?soComNota=true")
    assert so_nota.status_code == 200, so_nota.text
    aba_nota = load_workbook(BytesIO(so_nota.content)).active
    linhas_nota = list(aba_nota.iter_rows(min_row=2, values_only=True))
    assert [linha[0] for linha in linhas_nota] == [1]
    assert linhas_nota[0][1] == "Fulana Costa"

    triagem = client.get(f"/api/projetos/{pid}/triagem/exportar")
    assert triagem.status_code == 200, triagem.text
    aba_triagem = load_workbook(BytesIO(triagem.content)).active
    assert aba_triagem.title == "Triagem"
    cabecalho = [celula.value for celula in aba_triagem[1]]
    assert cabecalho == COLUNAS_TRIAGEM
    assert "Decisão" in cabecalho
    assert "Motivos" in cabecalho
    assert "Nota técnica" not in cabecalho


def test_banco_antigo_ganha_colunas_de_ranking(tmp_path):
    from persistencia.db import _garantir_colunas_rankings

    engine = create_engine(f"sqlite:///{(tmp_path / 'old.sqlite').as_posix()}")
    with engine.begin() as conn:
        conn.exec_driver_sql("CREATE TABLE avaliacoes_ia (id INTEGER PRIMARY KEY)")
    _garantir_colunas_rankings(engine)
    _garantir_colunas_rankings(engine)
    with engine.connect() as conn:
        colunas = {row[1] for row in conn.exec_driver_sql("PRAGMA table_info(avaliacoes_ia)")}
    assert {"nota_final", "nota_tecnica_zerada_por_falta_de_informacao"} <= colunas
