"""CA-JN-01 a CA-JN-09 da justificativa da nota técnica e da nota comportamental."""

from __future__ import annotations

import json
from io import BytesIO

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from avaliacao_ia.exportar import tabela_exportacao
from avaliacao_ia.prompt import PARAGRAFO_JUSTIFICATIVAS_NOTA, TEMPLATE_INICIAL
from avaliacao_ia.validacao import FRASE_ZERAMENTO_TECNICO, ErroValidacaoJson, validar_saida_modelo
from tests.test_rankings_ca import (
    PERGUNTA,
    RESPOSTA,
    _avaliar,
    _corpo,
    _ler,
    _ligar,
    _por_nome,
    _preparar,
)
from triagem.exportar import COLUNAS as COLUNAS_TRIAGEM

@pytest.fixture(autouse=True)
def sem_chamada_externa(monkeypatch):
    class ClienteHttpBloqueado(httpx.Client):
        def post(self, *args, **kwargs):
            raise AssertionError("chamada externa")

    monkeypatch.setattr("avaliacao_ia.cliente.httpx.Client", ClienteHttpBloqueado)

    def completar_proibido(self, mensagens):
        raise AssertionError("chamada externa")

    monkeypatch.setattr("avaliacao_ia.cliente.ClienteOpenAICompat.completar", completar_proibido)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("SCREENER_DB_PATH", str(tmp_path / "screener.sqlite"))
    import persistencia.db as banco

    banco._engine = None
    banco.init_db()
    from main import app

    with TestClient(app) as c:
        yield c


TECNICA = "A resposta cobre o pedido da vaga. Isso cabe na faixa da rubrica técnica."
COMPORTAMENTAL = "O texto é claro e organizado. Isso cabe na faixa da escrita."
PROMPT_CINCO = """Você avalia respostas de chatbot.
Não julgue cargo, idade nem escolaridade.
O JSON tem exatamente as chaves: perfil, por_pergunta, trechos, nota_tecnica, nota_comportamental.

## Job description
{{job_description}}

## Identificador
{{candidato_id}}

## Perguntas e respostas
{{perguntas_respostas}}
"""


def _respostas() -> list[dict]:
    return [{"colunaMapeada": PERGUNTA, "texto": RESPOSTA}]


def test_ca_jn_01_dois_textos_distintos_do_resumo(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    _ligar(monkeypatch, [_corpo(78, 70)])
    _avaliar(client, pid, [fulana["candidatoId"]])
    linha = _por_nome(_ler(client, pid))["Fulana Costa"]
    avaliacao = linha["avaliacao"]
    assert avaliacao["perfil"] == "A resposta foi lida contra a pergunta e a job description."
    assert avaliacao["justificativaNotaTecnica"] == TECNICA
    assert avaliacao["justificativaNotaComportamental"] == COMPORTAMENTAL
    assert avaliacao["perfil"] != avaliacao["justificativaNotaTecnica"]
    assert avaliacao["perfil"] != avaliacao["justificativaNotaComportamental"]


def test_ca_jn_02_justificativa_ausente_falha_depois_do_reparo(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    incompleto = json.loads(_corpo(80, 70))
    del incompleto["justificativa_nota_tecnica"]
    del incompleto["justificativa_nota_comportamental"]

    class Reparo:
        def __init__(self) -> None:
            self.mensagens: list[list[dict]] = []

        def completar(self, mensagens):
            self.mensagens.append(mensagens)
            return {
                "conteudo": json.dumps(incompleto),
                "uso": {"entrada": 1, "saida": 1, "total": 2, "nota": None},
                "tentativasRede": 1,
            }

    espiao = Reparo()
    monkeypatch.setattr("avaliacao_ia.servico.config.llm_configurada", lambda: True)
    monkeypatch.setattr("avaliacao_ia.servico.cliente_padrao", lambda: espiao)
    try:
        validar_saida_modelo(json.dumps(incompleto), _respostas())
        raise AssertionError("json sem justificativa_nota_tecnica deveria falhar")
    except ErroValidacaoJson as erro:
        assert "justificativa_nota_tecnica" in erro.mensagem

    saida = _avaliar(client, pid, [fulana["candidatoId"]])
    assert saida["resultados"][0]["estado"] == "falha_avaliacao"
    assert len(espiao.mensagens) == 2
    reparo = espiao.mensagens[1][-1]["content"]
    assert "justificativa_nota_tecnica" in reparo
    linha = _por_nome(_ler(client, pid))["Fulana Costa"]
    assert linha["resultadoExibicao"] == "falha_avaliacao"
    assert linha["decisao"]["resultado"] == "elegivel"
    assert linha["avaliacao"]["notaTecnica"] is None
    assert linha["avaliacao"]["notaComportamental"] is None
    assert linha["avaliacao"]["justificativaNotaTecnica"] is None
    assert linha["avaliacao"]["justificativaNotaComportamental"] is None
    assert linha["avaliacao"]["perfil"] is None
    assert linha["ranking"]["posicaoFinal"] is None


def test_ca_jn_03_string_vazia_cita_a_chave():
    corpo = json.loads(_corpo(70, 60))
    corpo["justificativa_nota_comportamental"] = "   "
    try:
        validar_saida_modelo(json.dumps(corpo), _respostas())
        raise AssertionError("justificativa em branco deveria falhar")
    except ErroValidacaoJson as erro:
        assert erro.codigo == "json_invalido"
        assert "justificativa_nota_comportamental" in erro.mensagem


def test_ca_jn_04_chave_extra_nao_gera_posicao(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    extra = _corpo(80, 70, extra={"ranking": 1})
    _ligar(monkeypatch, [extra, extra])
    saida = _avaliar(client, pid, [fulana["candidatoId"]])
    assert saida["resultados"][0]["estado"] == "falha_avaliacao"
    linha = _por_nome(_ler(client, pid))["Fulana Costa"]
    assert linha["ranking"]["posicaoFinal"] is None
    assert linha["ranking"]["posicaoTecnica"] is None
    assert linha["avaliacao"]["notaTecnica"] is None


def test_ca_jn_05_zeramento_tecnico_troca_a_justificativa(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    _ligar(monkeypatch, [_corpo(90, 40, resultado="sem_informacao")])
    _avaliar(client, pid, [fulana["candidatoId"]])
    linha = _por_nome(_ler(client, pid))["Fulana Costa"]
    assert linha["avaliacao"]["notaTecnica"] == 0
    assert linha["avaliacao"]["justificativaNotaTecnica"] == FRASE_ZERAMENTO_TECNICO
    assert linha["avaliacao"]["justificativaNotaComportamental"] == COMPORTAMENTAL
    assert linha["ranking"]["notaFinal"] == 20


def test_ca_jn_06_avaliacao_antiga_exporta_celulas_vazias(client: TestClient):
    from persistencia.db import SessionLocal
    from persistencia.models import AvaliacaoIa

    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    db = SessionLocal()
    try:
        db.add(
            AvaliacaoIa(
                candidato_id=fulana["candidatoId"],
                projeto_id=pid,
                prompt_version_id="prm_antigo",
                criterios_qualitativos_versao_id="",
                perfil="Resumo já gravado.",
                por_pergunta=[
                    {
                        "perguntaId": "p1",
                        "pergunta": PERGUNTA,
                        "resultado": "atende",
                        "justificativa": "A resposta conversa com a pergunta.",
                    }
                ],
                trechos=[{"perguntaId": "p1", "trecho": RESPOSTA}],
                nota_tecnica=78,
                nota_comportamental=70,
                nota_final=74,
                nota_tecnica_zerada_por_falta_de_informacao=False,
                justificativa_nota_tecnica=None,
                justificativa_nota_comportamental=None,
                estado="sucesso",
                erro=None,
            )
        )
        db.commit()
    finally:
        db.close()

    linha = _por_nome(_ler(client, pid))["Fulana Costa"]
    assert linha["avaliacao"]["notaTecnica"] == 78
    assert linha["avaliacao"]["notaComportamental"] == 70
    assert linha["avaliacao"]["perfil"] == "Resumo já gravado."
    assert linha["ranking"]["posicaoFinal"] == 1
    assert linha["avaliacao"]["justificativaNotaTecnica"] is None
    assert linha["avaliacao"]["justificativaNotaComportamental"] is None

    resultado = client.get(f"/api/projetos/{pid}/resultado/exportar")
    assert resultado.status_code == 200, resultado.text
    aba = load_workbook(BytesIO(resultado.content)).active
    cabecalho = [celula.value for celula in aba[1]]
    por_linha = {linha[0]: linha for linha in aba.iter_rows(min_row=2, values_only=True)}
    avaliada = por_linha[1]
    assert avaliada[cabecalho.index("Nota técnica")] == 78
    assert avaliada[cabecalho.index("Justificativa da nota técnica")] is None
    assert avaliada[cabecalho.index("Justificativa da nota comportamental")] is None
    assert avaliada[cabecalho.index("Resumo")] == "Resumo já gravado."
    assert "Busca LKD" not in cabecalho


def test_ca_jn_07_colunas_do_excel():
    por_pergunta = [
        {"pergunta": "P1", "resultado": "atende", "justificativa": "J1"},
        {"pergunta": "P2", "resultado": "atende", "justificativa": "J2"},
        {"pergunta": "P3", "resultado": "nao_atende", "justificativa": "J3"},
    ]
    linhas = [
        {
            "candidatoId": "c1",
            "linhaOrigem": 1,
            "ranking": {
                "posicaoFinal": 1,
                "posicaoTecnica": 1,
                "posicaoComportamental": 1,
                "notaFinal": 80,
            },
            "avaliacao": {
                "notaTecnica": 80,
                "notaComportamental": 70,
                "justificativaNotaTecnica": TECNICA,
                "justificativaNotaComportamental": COMPORTAMENTAL,
                "perfil": "Resumo.",
                "porPergunta": por_pergunta,
            },
        }
    ]
    cabecalho, tabela = tabela_exportacao(linhas, {"c1": {"nome": "Ana", "email": "a@x.com", "telefone": "1"}})
    trecho = cabecalho[
        cabecalho.index("Nota técnica") : cabecalho.index("Ranking comportamental") + 1
    ]
    assert trecho == [
        "Nota técnica",
        "Justificativa da nota técnica",
        "Ranking técnico",
        "Nota comportamental",
        "Justificativa da nota comportamental",
        "Ranking comportamental",
    ]
    depois = cabecalho[cabecalho.index("Resumo") + 1 :]
    assert depois == [
        "Justificativa pergunta 1",
        "Justificativa pergunta 2",
        "Justificativa pergunta 3",
    ]
    assert "Busca LKD" not in cabecalho
    assert "Busca LKD" not in COLUNAS_TRIAGEM
    assert tabela[0][cabecalho.index("Justificativa da nota técnica")] == TECNICA
    assert tabela[0][cabecalho.index("Justificativa pergunta 1")] == "J1"


def test_ca_jn_08_prompt_salvo_com_cinco_chaves_ainda_dispara(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    salvo = client.put(f"/api/projetos/{pid}/prompt", json={"template": PROMPT_CINCO})
    assert salvo.status_code == 200, salvo.text
    assert "justificativa_nota_tecnica" not in PROMPT_CINCO
    assert "7." in TEMPLATE_INICIAL
    espiao = _ligar(monkeypatch, [_corpo(78, 70)])
    preview = client.post(f"/api/projetos/{pid}/prompt/preview", json={"candidatoId": fulana["candidatoId"]})
    assert preview.status_code == 200, preview.text
    assert PARAGRAFO_JUSTIFICATIVAS_NOTA in preview.json()["texto"]
    assert espiao.chamadas == 0
    saida = _avaliar(client, pid, [fulana["candidatoId"]])
    assert saida["resultados"][0]["estado"] == "sucesso"
    assert espiao.chamadas == 1
    linha = _por_nome(_ler(client, pid))["Fulana Costa"]
    assert linha["avaliacao"]["justificativaNotaTecnica"] == TECNICA
    assert linha["avaliacao"]["notaTecnica"] == 78


class _FilaComMensagens:
    def __init__(self, corpos: list[str]) -> None:
        self.corpos = corpos
        self.chamadas = 0
        self.mensagens: list[list[dict]] = []

    def completar(self, mensagens):
        self.mensagens.append(mensagens)
        corpo = self.corpos[self.chamadas]
        self.chamadas += 1
        return {
            "conteudo": corpo,
            "uso": {"entrada": 4, "saida": 3, "total": 7, "nota": None},
            "tentativasRede": 1,
        }


def test_ca_jn_08_chamada_acrescenta_o_paragrafo(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    client.put(f"/api/projetos/{pid}/prompt", json={"template": PROMPT_CINCO})
    fila = _FilaComMensagens([_corpo(78, 70)])
    monkeypatch.setattr("avaliacao_ia.servico.config.llm_configurada", lambda: True)
    monkeypatch.setattr("avaliacao_ia.servico.cliente_padrao", lambda: fila)
    _avaliar(client, pid, [fulana["candidatoId"]])
    usuario = fila.mensagens[0][1]["content"]
    assert PARAGRAFO_JUSTIFICATIVAS_NOTA in usuario
    assert "justificativa_nota_tecnica" in fila.mensagens[0][0]["content"]


def test_ca_jn_09_so_quem_tem_nota(client: TestClient, monkeypatch):
    pid, pessoas = _preparar(client)
    fulana = pessoas["Fulana Costa"]
    _ligar(monkeypatch, [_corpo(78, 70)])
    _avaliar(client, pid, [fulana["candidatoId"]])
    resultado = client.get(f"/api/projetos/{pid}/resultado/exportar?soComNota=true")
    assert resultado.status_code == 200, resultado.text
    aba = load_workbook(BytesIO(resultado.content)).active
    cabecalho = [celula.value for celula in aba[1]]
    linhas = list(aba.iter_rows(min_row=2, values_only=True))
    assert [linha[1] for linha in linhas] == ["Fulana Costa"]
    assert "Justificativa da nota técnica" in cabecalho
    assert "Justificativa da nota comportamental" in cabecalho
    assert "Busca LKD" not in cabecalho


def test_trecho_em_texto_solto_vale_como_objeto():
    corpo = json.loads(_corpo(78, 70))
    corpo["trechos"] = [RESPOSTA]
    validado = validar_saida_modelo(json.dumps(corpo), _respostas())
    assert validado["trechos"] == [{"perguntaId": "p1", "trecho": RESPOSTA}]


def test_banco_antigo_ganha_colunas_de_justificativa(tmp_path):
    from sqlalchemy import create_engine

    from persistencia.db import _garantir_colunas_justificativa_notas

    engine = create_engine(f"sqlite:///{(tmp_path / 'old.sqlite').as_posix()}")
    with engine.begin() as conn:
        conn.exec_driver_sql("CREATE TABLE avaliacoes_ia (id INTEGER PRIMARY KEY, nota_tecnica INTEGER)")
    _garantir_colunas_justificativa_notas(engine)
    _garantir_colunas_justificativa_notas(engine)
    with engine.connect() as conn:
        colunas = {row[1] for row in conn.exec_driver_sql("PRAGMA table_info(avaliacoes_ia)")}
    assert {"justificativa_nota_tecnica", "justificativa_nota_comportamental"} <= colunas
