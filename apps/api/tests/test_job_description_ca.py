"""CA-JD-01 a CA-JD-08 da job description."""

from __future__ import annotations

import io
import json

import httpx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, func, select

from avaliacao_ia.prompt import MSG_JOB_DESCRIPTION_DESALINHADA
from job_description.extrair import MSG_DOCX_SEM_TEXTO, MSG_EXTENSAO, MSG_UTF8, MSG_VAZIA

TEXTO_V1 = "Descricao vigente numero um da vaga de analista."
TEXTO_V2 = "Descricao vigente numero dois da vaga revisada."
RESPOSTA_ELEGIVEL = "Disponibilidade integral de segunda a sexta no horario comercial."
RESPOSTA_INELEGIVEL = "Texto de resposta utilizavel do candidato inelegivel."
FRASE_PREFIXO = (
    "O texto marcado como job description é a descrição geral da vaga. "
    "Use-o como contexto do que a vaga pede. Ele não reabre cargo, idade nem escolaridade."
)
DADOS_FORA = (
    "Fulana de Tal",
    "Bacharelado",
    "Recife",
    "status-legado-xyz",
    "Ciclano Souza",
    "Operador de caixa",
    "Ensino medio",
    "Olinda",
)


@pytest.fixture(autouse=True)
def sem_chamada_externa(monkeypatch):
    class ClienteHttpBloqueado(httpx.Client):
        def post(self, *args, **kwargs):
            raise AssertionError("chamada externa")

    monkeypatch.setattr("avaliacao_ia.cliente.httpx.Client", ClienteHttpBloqueado)

    def completar_proibido(self, mensagens):
        raise AssertionError("chamada externa")

    monkeypatch.setattr("avaliacao_ia.cliente.ClienteOpenAICompat.completar", completar_proibido)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("SCREENER_DB_PATH", str(tmp_path / "screener.sqlite"))
    import persistencia.db as banco

    banco._engine = None
    banco.init_db()
    from main import app

    with TestClient(app) as c:
        yield c


def _criar(client: TestClient, nome: str = "Vaga") -> str:
    r = client.post("/api/projetos", json={"nome": nome})
    assert r.status_code == 200, r.text
    return r.json()["projetoId"]


def _criterios_e_lote(client: TestClient, pid: str) -> tuple[dict, dict]:
    r = client.put(
        f"/api/projetos/{pid}/criterios-deterministicos",
        json={
            "cargos": {"ativo": True, "aceitos": ["Analista de dados"]},
            "idade": {"ativo": False, "minima": None, "maxima": None},
            "escolaridade": {
                "ativo": False,
                "configuracao": {"modelo": "lista_valores_aceitos", "aceitos": []},
            },
        },
    )
    assert r.status_code == 200, r.text
    csv = (
        "Nome,Cargo,Idade,Escolaridade,Cidade,Status,Resposta\n"
        f"Fulana de Tal,Analista de dados,34,Bacharelado,Recife,status-legado-xyz,{RESPOSTA_ELEGIVEL}\n"
        f"Ciclano Souza,Operador de caixa,41,Ensino medio,Olinda,status-legado-xyz,{RESPOSTA_INELEGIVEL}\n"
    )
    mapa = {
        "ultimoCargo": "Cargo",
        "idade": "Idade",
        "escolaridade": "Escolaridade",
        "nomeExibicao": "Nome",
        "respostasChatbot": ["Resposta"],
        "cidade": "Cidade",
        "statusLegado": "Status",
    }
    r = client.post(
        f"/api/projetos/{pid}/importacao/confirmar",
        files={"arquivo": ("lote.csv", csv.encode("utf-8"), "text/csv")},
        data={"encoding": "utf-8", "mapeamento_json": json.dumps(mapa)},
    )
    assert r.status_code == 200, r.text
    r = client.post(f"/api/projetos/{pid}/triagem/executar")
    assert r.status_code == 200, r.text
    linhas = client.get(f"/api/projetos/{pid}/triagem").json()["linhas"]
    eleg = next(l for l in linhas if l["decisao"]["resultado"] == "elegivel")
    ineleg = next(l for l in linhas if l["decisao"]["resultado"] == "inelegivel")
    return eleg, ineleg


def _qualificar(client: TestClient, pid: str) -> dict:
    r = client.put(
        f"/api/projetos/{pid}/criterios-qualitativos",
        json={
            "itens": [
                {
                    "nome": "Turnos de trabalho",
                    "descricao": "Horarios informados na resposta",
                    "ordem": 1,
                    "eliminatorio": False,
                }
            ]
        },
    )
    assert r.status_code == 200, r.text
    return r.json()["itens"][0]


def _salvar_prompt(client: TestClient, pid: str) -> dict:
    inicial = client.get(f"/api/projetos/{pid}/prompt").json()["templateInicial"]
    assert "{{job_description}}" in inicial
    assert "{{perguntas_respostas}}" in inicial
    assert "{{candidato_id}}" in inicial
    assert "{{criterios_qualitativos}}" not in inicial
    r = client.put(f"/api/projetos/{pid}/prompt", json={"template": inicial})
    assert r.status_code == 200, r.text
    return r.json()


def _salvar_jd(client: TestClient, pid: str, texto: str) -> dict:
    r = client.put(f"/api/projetos/{pid}/job-description", json={"texto": texto})
    assert r.status_code == 200, r.text
    return r.json()


def _json_modelo(trecho: str) -> str:
    return json.dumps(
        {
            "perfil": "A pessoa informa disponibilidade de horario comercial.",
            "por_pergunta": [
                {
                    "perguntaId": "p1",
                    "pergunta": "Resposta",
                    "resultado": "atende",
                    "justificativa": "A resposta informa o horario, em linha com a pergunta e a job description.",
                }
            ],
            "trechos": [{"perguntaId": "p1", "trecho": trecho}],
            "nota_tecnica": 70,
            "nota_comportamental": 65,
            "justificativa_nota_tecnica": "A resposta cobre o pedido da vaga. Isso cabe na faixa da rubrica técnica.",
            "justificativa_nota_comportamental": "O texto é claro e organizado. Isso cabe na faixa da escrita.",
        }
    )


def _ligar_modelo(monkeypatch, conteudo: str):
    class Espiao:
        def __init__(self) -> None:
            self.chamadas = 0
            self.mensagens: list[list[dict]] = []

        def completar(self, mensagens):
            self.chamadas += 1
            self.mensagens.append(mensagens)
            return {
                "conteudo": conteudo,
                "uso": {"entrada": 3, "saida": 2, "total": 5, "nota": None},
                "tentativasRede": 1,
            }

    espiao = Espiao()
    monkeypatch.setattr("avaliacao_ia.servico.config.llm_configurada", lambda: True)
    monkeypatch.setattr("avaliacao_ia.servico.cliente_padrao", lambda: espiao)
    return espiao


def _avaliacao_de(client: TestClient, pid: str, candidato_id: str) -> dict:
    linhas = client.get(f"/api/projetos/{pid}/resultado").json()["linhas"]
    linha = next(l for l in linhas if l["candidatoId"] == candidato_id)
    assert linha["avaliacao"] is not None
    return linha["avaliacao"]


def test_ca_jd_01_salvar_texto_sem_chamada_externa(client: TestClient):
    pid = _criar(client)
    bruto = "  Texto da vaga de analista.\nCom segunda linha.  "
    up = client.post(
        f"/api/projetos/{pid}/job-description/arquivo",
        files={"arquivo": ("vaga.txt", bruto.encode("utf-8"), "text/plain")},
        data={"encoding": "utf-8"},
    )
    assert up.status_code == 200, up.text
    assert up.json()["texto"] == bruto
    assert client.get(f"/api/projetos/{pid}/job-description").json()["vigente"] is None

    salvo = client.put(f"/api/projetos/{pid}/job-description", json={"texto": up.json()["texto"]})
    assert salvo.status_code == 200, salvo.text
    corpo = salvo.json()
    assert corpo["versao"] == 1
    assert corpo["texto"] == bruto.strip()
    assert corpo["jobDescriptionId"]
    got = client.get(f"/api/projetos/{pid}/job-description").json()
    assert got["vigente"]["jobDescriptionId"] == corpo["jobDescriptionId"]
    assert got["vigente"]["texto"] == bruto.strip()
    assert got["versoes"] == [got["vigente"]]


def test_ca_jd_02_vazio_nao_grava(client: TestClient):
    pid = _criar(client)
    vazio = client.put(f"/api/projetos/{pid}/job-description", json={"texto": " \n\t "})
    assert vazio.status_code == 400
    assert vazio.json()["detail"] == MSG_VAZIA
    assert client.get(f"/api/projetos/{pid}/job-description").json()["versoes"] == []

    primeiro = _salvar_jd(client, pid, TEXTO_V1)
    de_novo = client.put(f"/api/projetos/{pid}/job-description", json={"texto": "   "})
    assert de_novo.status_code == 400
    assert de_novo.json()["detail"] == MSG_VAZIA
    got = client.get(f"/api/projetos/{pid}/job-description").json()
    assert len(got["versoes"]) == 1
    assert got["vigente"]["jobDescriptionId"] == primeiro["jobDescriptionId"]
    assert got["vigente"]["texto"] == TEXTO_V1


def test_ca_jd_03_pdf_recusado_nao_muda_vigente(client: TestClient):
    pid = _criar(client)
    primeiro = _salvar_jd(client, pid, TEXTO_V1)
    pdf = client.post(
        f"/api/projetos/{pid}/job-description/arquivo",
        files={"arquivo": ("vaga.pdf", b"%PDF-1.4 fake", "application/pdf")},
        data={"encoding": "utf-8"},
    )
    assert pdf.status_code == 400
    assert pdf.json()["detail"] == MSG_EXTENSAO
    outro = client.post(
        f"/api/projetos/{pid}/job-description/arquivo",
        files={"arquivo": ("vaga.xlsx", b"PK fake", "application/octet-stream")},
        data={"encoding": "utf-8"},
    )
    assert outro.status_code == 400
    assert outro.json()["detail"] == MSG_EXTENSAO
    got = client.get(f"/api/projetos/{pid}/job-description").json()
    assert got["vigente"]["jobDescriptionId"] == primeiro["jobDescriptionId"]
    assert got["vigente"]["texto"] == TEXTO_V1
    assert len(got["versoes"]) == 1


def test_ca_jd_04_triagem_sem_job_description(client: TestClient):
    pid = _criar(client)
    assert client.get(f"/api/projetos/{pid}/job-description").json()["vigente"] is None
    _criterios_e_lote(client, pid)
    r = client.get(f"/api/projetos/{pid}/triagem")
    assert r.status_code == 200
    linhas = r.json()["linhas"]
    assert len(linhas) == 2
    por_resultado = {l["decisao"]["resultado"] for l in linhas}
    assert por_resultado == {"elegivel", "inelegivel"}
    ineleg = next(l for l in linhas if l["decisao"]["resultado"] == "inelegivel")
    assert ineleg["decisao"]["motivos"][0]["codigo"] == "cargo_fora_da_lista"
    bruto = json.dumps(r.json()).lower()
    assert "job description" not in bruto
    assert "job_description" not in bruto


def test_ca_jd_05_ia_bloqueada_sem_descricao_vigente_no_prompt(client: TestClient, monkeypatch):
    pid = _criar(client)
    eleg, _ineleg = _criterios_e_lote(client, pid)
    _qualificar(client, pid)
    v1 = _salvar_jd(client, pid, TEXTO_V1)
    prompt = _salvar_prompt(client, pid)
    assert prompt["jobDescriptionVersaoId"] == v1["jobDescriptionId"]
    assert prompt["jobDescriptionSnapshot"] == TEXTO_V1
    v2 = _salvar_jd(client, pid, TEXTO_V2)
    assert v2["jobDescriptionId"] != v1["jobDescriptionId"]

    aviso = client.get(f"/api/projetos/{pid}/job-description").json()["aviso"]
    assert aviso == MSG_JOB_DESCRIPTION_DESALINHADA
    gravado = client.get(f"/api/projetos/{pid}/prompt").json()["vigente"]
    assert gravado["jobDescriptionSnapshot"] == TEXTO_V1
    assert gravado["jobDescriptionVersaoId"] == v1["jobDescriptionId"]

    espiao = _ligar_modelo(monkeypatch, "{}")
    preview = client.post(f"/api/projetos/{pid}/prompt/preview", json={"candidatoId": eleg["candidatoId"]})
    assert preview.status_code == 200, preview.text
    texto = preview.json()["texto"]
    assert TEXTO_V2 in texto
    assert TEXTO_V1 not in texto
    assert FRASE_PREFIXO in texto
    assert espiao.chamadas == 0

    disparo = client.post(
        f"/api/projetos/{pid}/avaliacao-ia/executar",
        json={"candidatoIds": [eleg["candidatoId"]], "confirmar": True},
    )
    assert disparo.status_code == 400
    assert disparo.json()["detail"] == MSG_JOB_DESCRIPTION_DESALINHADA
    assert espiao.chamadas == 0


def test_ca_jd_06_inelegivel_nao_recebe_o_texto(client: TestClient, monkeypatch):
    pid = _criar(client)
    eleg, ineleg = _criterios_e_lote(client, pid)
    item = _qualificar(client, pid)
    _salvar_jd(client, pid, TEXTO_V1)
    _salvar_prompt(client, pid)
    trecho = eleg["respostasChatbot"][0]["texto"]
    espiao = _ligar_modelo(monkeypatch, _json_modelo(trecho))
    disparo = client.post(
        f"/api/projetos/{pid}/avaliacao-ia/executar",
        json={"candidatoIds": [eleg["candidatoId"], ineleg["candidatoId"]], "confirmar": True},
    )
    assert disparo.status_code == 200, disparo.text
    assert disparo.json()["enviados"] == [eleg["candidatoId"]]
    assert espiao.chamadas == 1
    payload = "\n".join(m["content"] for m in espiao.mensagens[0])
    assert TEXTO_V1 in payload
    assert trecho in payload
    assert eleg["candidatoId"] in payload
    assert ineleg["candidatoId"] not in payload
    assert RESPOSTA_INELEGIVEL not in payload
    sem_ids = payload.replace(eleg["candidatoId"], "").replace(item["criterioId"], "")
    assert "34" not in sem_ids
    assert "41" not in sem_ids
    for dado in (*DADOS_FORA, "Analista de dados"):
        assert dado not in payload


def test_ca_jd_07_duplicar_copia_descricao_e_nao_o_lote(client: TestClient, monkeypatch):
    pid = _criar(client, "Projeto A")
    eleg, _ineleg = _criterios_e_lote(client, pid)
    _qualificar(client, pid)
    v1 = _salvar_jd(client, pid, TEXTO_V1)
    _salvar_prompt(client, pid)
    trecho = eleg["respostasChatbot"][0]["texto"]
    _ligar_modelo(monkeypatch, _json_modelo(trecho))
    disparo = client.post(
        f"/api/projetos/{pid}/avaliacao-ia/executar",
        json={"candidatoIds": [eleg["candidatoId"]], "confirmar": True},
    )
    assert disparo.status_code == 200, disparo.text

    copia = client.post(f"/api/projetos/{pid}/duplicar", json={"nome": "Projeto B"})
    assert copia.status_code == 200, copia.text
    novo = copia.json()["projetoId"]
    assert novo != pid
    assert copia.json()["quantidadeCandidatos"] == 0
    assert copia.json()["promptVigenteId"] is None
    assert copia.json()["jobDescriptionVigenteId"] != v1["jobDescriptionId"]
    jd = client.get(f"/api/projetos/{novo}/job-description").json()
    assert jd["vigente"]["texto"] == TEXTO_V1
    assert jd["vigente"]["versao"] == 1
    assert jd["vigente"]["jobDescriptionId"] != v1["jobDescriptionId"]
    assert client.get(f"/api/projetos/{novo}/prompt").json()["vigente"] is None
    assert client.get(f"/api/projetos/{novo}/triagem").json()["linhas"] == []

    from persistencia.db import SessionLocal
    from persistencia.models import AvaliacaoIa

    db = SessionLocal()
    try:
        origem = db.scalar(
            select(func.count()).select_from(AvaliacaoIa).where(AvaliacaoIa.projeto_id == pid)
        )
        destino = db.scalar(
            select(func.count()).select_from(AvaliacaoIa).where(AvaliacaoIa.projeto_id == novo)
        )
    finally:
        db.close()
    assert origem == 1
    assert destino == 0


def test_ca_jd_08_versao_usada_nao_muda(client: TestClient, monkeypatch):
    pid = _criar(client)
    eleg, _ineleg = _criterios_e_lote(client, pid)
    _qualificar(client, pid)
    v1 = _salvar_jd(client, pid, TEXTO_V1)
    _salvar_prompt(client, pid)
    trecho = eleg["respostasChatbot"][0]["texto"]
    _ligar_modelo(monkeypatch, _json_modelo(trecho))
    disparo = client.post(
        f"/api/projetos/{pid}/avaliacao-ia/executar",
        json={"candidatoIds": [eleg["candidatoId"]], "confirmar": True},
    )
    assert disparo.status_code == 200, disparo.text
    avaliacao = _avaliacao_de(client, pid, eleg["candidatoId"])
    assert avaliacao["jobDescriptionVersaoId"] == v1["jobDescriptionId"]
    assert avaliacao["estado"] == "sucesso"

    v2 = _salvar_jd(client, pid, TEXTO_V2)
    assert v2["jobDescriptionId"] != v1["jobDescriptionId"]
    de_novo = _avaliacao_de(client, pid, eleg["candidatoId"])
    assert de_novo["jobDescriptionVersaoId"] == v1["jobDescriptionId"]
    versoes = client.get(f"/api/projetos/{pid}/job-description").json()["versoes"]
    assert [v["texto"] for v in versoes] == [TEXTO_V1, TEXTO_V2]
    assert versoes[0]["jobDescriptionId"] == v1["jobDescriptionId"]


def test_txt_utf8_invalido_pede_latin1_e_nao_grava(client: TestClient):
    pid = _criar(client)
    _salvar_jd(client, pid, TEXTO_V1)
    conteudo = "Descrição com acento".encode("latin-1")
    ruim = client.post(
        f"/api/projetos/{pid}/job-description/arquivo",
        files={"arquivo": ("vaga.txt", conteudo, "text/plain")},
        data={"encoding": "utf-8"},
    )
    assert ruim.status_code == 400
    assert ruim.json()["detail"] == MSG_UTF8
    bom = client.post(
        f"/api/projetos/{pid}/job-description/arquivo",
        files={"arquivo": ("vaga.txt", conteudo, "text/plain")},
        data={"encoding": "latin-1"},
    )
    assert bom.status_code == 200, bom.text
    assert "Descrição" in bom.json()["texto"]
    assert client.get(f"/api/projetos/{pid}/job-description").json()["vigente"]["texto"] == TEXTO_V1


def test_docx_extrai_paragrafos_sem_gravar(client: TestClient):
    from docx import Document

    pid = _criar(client)
    doc = Document()
    doc.add_paragraph("Primeiro paragrafo da vaga.")
    doc.add_paragraph("Segundo paragrafo da vaga.")
    buf = io.BytesIO()
    doc.save(buf)
    up = client.post(
        f"/api/projetos/{pid}/job-description/arquivo",
        files={
            "arquivo": (
                "vaga.docx",
                buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"encoding": "latin-1"},
    )
    assert up.status_code == 200, up.text
    texto = up.json()["texto"]
    assert texto.index("Primeiro paragrafo da vaga.") < texto.index("Segundo paragrafo da vaga.")
    assert client.get(f"/api/projetos/{pid}/job-description").json()["vigente"] is None

    vazio = Document()
    buf_vazio = io.BytesIO()
    vazio.save(buf_vazio)
    recusa = client.post(
        f"/api/projetos/{pid}/job-description/arquivo",
        files={
            "arquivo": (
                "vazio.docx",
                buf_vazio.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert recusa.status_code == 400
    assert recusa.json()["detail"] == MSG_DOCX_SEM_TEXTO


def test_banco_antigo_ganha_colunas_de_job_description(tmp_path):
    from persistencia.db import _garantir_colunas_job_description

    engine = create_engine(f"sqlite:///{(tmp_path / 'old.sqlite').as_posix()}")
    with engine.begin() as conn:
        conn.exec_driver_sql("CREATE TABLE projetos (id VARCHAR(64) PRIMARY KEY, nome VARCHAR(255) NOT NULL)")
        conn.exec_driver_sql("CREATE TABLE prompts (prompt_id VARCHAR(64) PRIMARY KEY, template TEXT)")
        conn.exec_driver_sql("CREATE TABLE avaliacoes_ia (id INTEGER PRIMARY KEY)")
    _garantir_colunas_job_description(engine)
    _garantir_colunas_job_description(engine)
    with engine.connect() as conn:
        projetos = {row[1] for row in conn.exec_driver_sql("PRAGMA table_info(projetos)")}
        prompts = {row[1] for row in conn.exec_driver_sql("PRAGMA table_info(prompts)")}
        avals = {row[1] for row in conn.exec_driver_sql("PRAGMA table_info(avaliacoes_ia)")}
    assert "job_description_vigente_id" in projetos
    assert "job_description_versao_id" in prompts
    assert "job_description_snapshot" in prompts
    assert "job_description_versao_id" in avals
